"""Excel workbook with live formulas for the photovoltaic and wind designs.

Qt-free. The workbook lets a student see every step of the calculation as a spreadsheet
formula, change an input, and watch the results move. Inputs are named cells, so the formulas
read like the textbook ones (``=coverage*demand_kwh/(irradiation*pr_effective)``).

The yearly energy models are reproduced exactly (PV: Hs x PR; wind: Weibull x power curve).
What is *not* reproduced is stated on the 'About' sheet (e.g. the wind-farm wake model).
A 'Check' sheet compares every key result with the value computed by RESize.
"""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from resize_energy.core.plant_economics import EconomicInputs
from resize_energy.core.pv_design import PVDesign, PVInputs
from resize_energy.core.wind_design import WindDesign, WindInputs

_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # yellow: you can change this
_HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
_BOLD = Font(bold=True)
_LAYOUT_ROWS = 100      # candidate string lengths examined in the PV layout table
_CASH_YEARS = 60        # rows of the cash-flow table (rows beyond the lifetime stay empty)


class _Book:
    """Small helper that writes labelled rows and registers a name for each value."""

    def __init__(self) -> None:
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.checks: list[tuple[str, str, float, str]] = []   # label, name, app value, format

    def sheet(self, title: str, widths: tuple[int, ...] = (44, 16, 14, 70)) -> Worksheet:
        ws = self.wb.create_sheet(title)
        for col, w in zip("ABCD", widths, strict=False):
            ws.column_dimensions[col].width = w
        return ws

    def heading(self, ws: Worksheet, text: str) -> None:
        ws.append([])
        ws.append([text])
        ws.cell(ws.max_row, 1).font = _BOLD
        for c in range(1, 5):
            ws.cell(ws.max_row, c).fill = _HEAD_FILL

    def put(self, ws: Worksheet, name: str, label: str, value, unit: str = "", note: str = "",
            fmt: str | None = None, is_input: bool = False, check: float | None = None) -> None:
        """One row: label | value or formula | unit | note. ``name`` is usable in formulas."""
        ws.append([label, value, unit, note])
        row = ws.max_row
        cell = ws.cell(row, 2)
        if fmt:
            cell.number_format = fmt
        if is_input:
            cell.fill = _INPUT_FILL
        ws.cell(row, 4).alignment = Alignment(wrap_text=True, vertical="top")
        if name:
            self.wb.defined_names[name] = DefinedName(
                name, attr_text=f"'{ws.title}'!$B${row}")
        if check is not None and math.isfinite(check):
            self.checks.append((label, name, float(check), fmt or "General"))

    def finish_checks(self) -> None:
        ws = self.sheet("Check", (44, 18, 18, 14))
        ws.append(["Result", "In this workbook", "Computed by RESize", "Status"])
        for c in range(1, 5):
            ws.cell(1, c).font = _BOLD
            ws.cell(1, c).fill = _HEAD_FILL
        for label, name, app_value, fmt in self.checks:
            ws.append([label, f"={name}", app_value, None])
            r = ws.max_row
            ws.cell(r, 2).number_format = fmt
            ws.cell(r, 3).number_format = fmt
            ws.cell(r, 4).value = (
                f'=IF(ABS(B{r}-C{r})<=0.0005*MAX(1,ABS(C{r})),"OK","DIFFERENT")')
        ws.append([])
        ws.append(["If you change an input on the other sheets, 'DIFFERENT' is expected: the "
                   "right-hand column keeps the values RESize computed for the original inputs."])
        ws.freeze_panes = "A2"


# --- the economic part, shared by PV and wind ----------------------------------------------
def _economics(book: _Book, e: EconomicInputs, power_name: str, energy_name: str,
               demand_name: str, app, power_label: str) -> None:
    ws = book.sheet("Economics")
    ws.append(["Costs, savings and carbon (textbook annuity model, constant prices)"])
    ws.cell(1, 1).font = _BOLD
    book.heading(ws, "Inputs (yellow cells can be changed)")
    p = book.put
    p(ws, "capex_per_kw", f"Plant cost per {power_label}", e.capex_per_kw, f"euro/{power_label}",
      "Placeholder: replace with a real quote.", "#,##0.00", True)
    p(ws, "om_per_kw_year", f"Operation and maintenance per {power_label} and year",
      e.om_per_kw_year, f"euro/{power_label}/year", "", "#,##0.00", True)
    p(ws, "life_years", "Lifetime", e.life_years, "years", "", "0.0", True)
    p(ws, "discount_rate", "Discount rate", e.discount_rate, "-", "Real rate, e.g. 0.05 = 5 %.",
      "0.0%", True)
    p(ws, "retail_price", "Price you pay for electricity", e.retail_price, "euro/kWh", "",
      "0.000", True)
    p(ws, "export_price", "Price paid for the surplus", e.export_price, "euro/kWh", "", "0.000",
      True)
    p(ws, "grid_ef", "Emission factor of the grid", e.grid_ef, "kg CO2/kWh", "", "0.000", True)
    p(ws, "lifecycle_g_per_kwh", "Life-cycle emissions of the plant", e.lifecycle_g_per_kwh,
      "g CO2/kWh", "", "0.0", True)
    book.heading(ws, "Calculation")
    ec = app
    p(ws, "capex", "Plant cost (CAPEX)", f"={power_name}*capex_per_kw", "euro",
      "= power x cost per unit of power", "#,##0", check=ec.capex)
    p(ws, "om_per_year", "O&M per year", f"={power_name}*om_per_kw_year", "euro/year", "",
      "#,##0", check=ec.om_per_year)
    p(ws, "self_used_kwh", "Energy used on site", f"=MIN({energy_name},{demand_name})", "kWh/year",
      "Annual balance: up to the demand.", "#,##0", check=ec.self_used_kwh)
    p(ws, "exported_kwh", "Energy exported", f"=MAX(0,{energy_name}-{demand_name})", "kWh/year",
      "", "#,##0", check=ec.exported_kwh)
    p(ws, "savings_per_year", "Value of the energy per year",
      "=self_used_kwh*retail_price+exported_kwh*export_price", "euro/year", "", "#,##0",
      check=ec.savings_per_year)
    p(ws, "net_benefit", "Net benefit per year", "=savings_per_year-om_per_year", "euro/year",
      "Savings minus O&M.", "#,##0", check=ec.net_benefit_per_year)
    p(ws, "crf", "Capital recovery factor",
      "=IF(discount_rate=0,1/life_years,discount_rate*(1+discount_rate)^life_years/"
      "((1+discount_rate)^life_years-1))", "1/year", "CRF = r(1+r)^n / ((1+r)^n - 1)", "0.0000")
    p(ws, "annuity", "Annuity factor", "=1/crf", "years",
      "Present value of 1 euro per year over the lifetime.", "0.000")
    p(ws, "lcoe", "Levelised cost of energy",
      f"=(capex*crf+om_per_year)/{energy_name}", "euro/kWh",
      "LCOE = (CAPEX x CRF + O&M) / yearly energy", "0.0000", check=ec.lcoe)
    p(ws, "payback", "Simple payback",
      '=IF(net_benefit>0,capex/net_benefit,"never")', "years", "CAPEX / net benefit", "0.0",
      check=ec.payback_years)
    p(ws, "disc_payback", "Discounted payback",
      '=IF(net_benefit<=0,"never",IF(discount_rate=0,capex/net_benefit,'
      'IF(capex*discount_rate/net_benefit<1,'
      '-LN(1-capex*discount_rate/net_benefit)/LN(1+discount_rate),"never")))', "years",
      "n = -ln(1 - CAPEX x r / net) / ln(1 + r)", "0.0", check=ec.discounted_payback_years)
    p(ws, "npv_euro", "Net present value", "=-capex+net_benefit*annuity", "euro",
      "NPV = -CAPEX + net benefit x annuity factor", "#,##0", check=ec.npv)
    p(ws, "co2_avoided", "CO2 avoided per year",
      f"={energy_name}*(grid_ef-lifecycle_g_per_kwh/1000)/1000", "t CO2/year",
      "Energy x (grid factor - plant life-cycle factor)", "#,##0.0",
      check=ec.co2_avoided_t_per_year)

    cf = book.sheet("Cash flow", (10, 22, 3, 3))
    cf.append(["Year", "Cumulative discounted result (euro)"])
    for c in (1, 2):
        cf.cell(1, c).font = _BOLD
        cf.cell(1, c).fill = _HEAD_FILL
    cf.append([0, "=-capex"])
    for y in range(1, _CASH_YEARS + 1):
        n = f"MIN(A{y + 2},life_years)"
        cf.append([y, f'=IF(A{y + 2}>ROUNDUP(life_years,0),"",-capex+net_benefit*'
                      f'IF(discount_rate=0,{n},(1-(1+discount_rate)^(-{n}))/discount_rate))'])
    for r in range(2, cf.max_row + 1):
        cf.cell(r, 2).number_format = "#,##0"
    cf.freeze_panes = "A2"


def _about(book: _Book, title: str, lines: list[str]) -> None:
    ws = book.wb.create_sheet("About", 0)
    ws.column_dimensions["A"].width = 110
    ws.append([title])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    for line in lines:
        ws.append([line])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")


_COMMON_ABOUT = [
    "Made by RESize. Yellow cells are inputs: change them and every result follows.",
    "Click a result and read its formula: the input names are the ones on the 'Inputs' sheet "
    "(Formulas > Name Manager lists them).",
    "The 'Check' sheet compares this workbook with the numbers RESize computed.",
    "Preliminary design only: confirm against datasheets, grid rules and local standards.",
]


# --- photovoltaic --------------------------------------------------------------------------
def build_pv_workbook(inp: PVInputs, design: PVDesign) -> Workbook:
    book = _Book()
    m, v = inp.module, inp.inverter
    _about(book, "RESize - photovoltaic design as a spreadsheet", _COMMON_ABOUT + [
        "Reproduced: target power, string limits from the inverter voltages, string layout "
        "(table on 'Layout'), yearly energy (irradiation x PR), economics.",
        "Not reproduced: the monthly profile (used only for charts). Half-way rounding in "
        "'Layout' can differ from RESize by one module in rare tie cases.",
    ])
    ws = book.sheet("Inputs")
    ws.append(["Inputs (yellow cells can be changed)"])
    ws.cell(1, 1).font = _BOLD
    p = book.put
    book.heading(ws, "Site and demand")
    p(ws, "demand_kwh", "Annual electricity demand", inp.annual_demand_kwh, "kWh/year", "", "#,##0",
      True)
    p(ws, "irradiation", "Yearly irradiation on the modules", inp.annual_irradiation_kwh_m2,
      "kWh/m2/year", "", "#,##0", True)
    p(ws, "coverage", "Share of the demand to cover", inp.coverage, "-", "1 = 100 %", "0.00", True)
    p(ws, "pr", "Performance ratio", inp.performance_ratio, "-", "", "0.00", True)
    p(ws, "shading", "Loss to shading", inp.shading_loss, "-", "", "0.00", True)
    p(ws, "t_min", "Lowest air temperature", inp.t_min_c, "degC", "", "0.0", True)
    p(ws, "t_hot", "Highest cell temperature", inp.t_cell_max_c, "degC", "", "0.0", True)
    book.heading(ws, "Module (datasheet, STC)")
    p(ws, "p_stc_w", "Power", m.p_stc_w, "W", m.name, "0.0", True)
    p(ws, "voc", "Open-circuit voltage", m.voc_v, "V", "", "0.00", True)
    p(ws, "vmp", "Voltage at maximum power", m.vmp_v, "V", "", "0.00", True)
    p(ws, "isc", "Short-circuit current", m.isc_a, "A", "", "0.00", True)
    p(ws, "tc_voc", "Temperature coefficient of Voc", m.temp_coeff_voc_pct_per_c, "%/degC", "",
      "0.000", True)
    p(ws, "tc_vmp", "Temperature coefficient of Vmp", m.temp_coeff_vmp_pct_per_c, "%/degC", "",
      "0.000", True)
    p(ws, "module_area", "Module area", m.area_m2, "m2", "", "0.00", True)
    book.heading(ws, "Inverter (datasheet)")
    p(ws, "p_ac_kw", "AC power", v.p_ac_kw, "kW", v.name, "0.00", True)
    p(ws, "vdc_max", "Maximum DC voltage", v.vdc_max_v, "V", "", "0", True)
    p(ws, "mppt_vmin", "MPPT window, minimum", v.mppt_vmin_v, "V", "", "0", True)
    p(ws, "mppt_vmax", "MPPT window, maximum", v.mppt_vmax_v, "V", "", "0", True)
    p(ws, "imax_mppt", "Maximum current per MPPT input", v.imax_per_mppt_a, "A", "", "0.0", True)
    p(ws, "n_mppt", "Number of MPPT inputs", v.n_mppt, "-", "", "0", True)

    calc = book.sheet("Calculation")
    calc.append(["Calculation (every cell is a formula)"])
    calc.cell(1, 1).font = _BOLD
    book.heading(calc, "1. How much power")
    p(calc, "pr_effective", "Effective performance ratio", "=pr*(1-shading)", "-",
      "PR x (1 - shading loss)", "0.000", check=design.performance_ratio_effective)
    p(calc, "target_kwp", "Target peak power", "=coverage*demand_kwh/(irradiation*pr_effective)",
      "kWp", "P = coverage x demand / (H x PR)", "0.000", check=design.target_kwp)
    p(calc, "modules_target", "Modules needed", "=MAX(1,ROUNDUP(target_kwp*1000/p_stc_w,0))",
      "modules", "Target power / module power, rounded up", "0", check=design.modules_target)
    book.heading(calc, "2. Voltage limits of a string")
    p(calc, "voc_cold", "Voc of one module at the lowest temperature",
      "=voc*(1+tc_voc/100*(t_min-25))", "V", "V(T) = V_stc x (1 + beta x (T - 25))", "0.00")
    p(calc, "vmp_cold", "Vmp of one module at the lowest temperature",
      "=vmp*(1+tc_vmp/100*(t_min-25))", "V", "", "0.00")
    p(calc, "vmp_hot", "Vmp of one module at the highest temperature",
      "=vmp*(1+tc_vmp/100*(t_hot-25))", "V", "", "0.00")
    p(calc, "n_max", "Longest string",
      "=MIN(INT(vdc_max/voc_cold),INT(mppt_vmax/vmp_cold))", "modules",
      "Cold Voc must stay below the inverter limit; cold Vmp below the MPPT maximum.", "0",
      check=design.n_series_max)
    p(calc, "n_min", "Shortest string", "=ROUNDUP(mppt_vmin/vmp_hot,0)", "modules",
      "Hot Vmp must stay above the MPPT minimum.", "0", check=design.n_series_min)
    p(calc, "feasible", "A valid string length exists", "=AND(n_min<=n_max,n_max>=1)", "", "")
    p(calc, "string_capacity", "Strings the inverter can take",
      "=MAX(1,INT(imax_mppt/isc))*n_mppt", "strings", "Current limit per MPPT input", "0",
      check=design.string_capacity)
    book.heading(calc, "3. String layout (see the 'Layout' sheet)")
    p(calc, "n_series", "Modules in series",
      "=IF(feasible,INDEX(Layout!A2:A101,MATCH(MIN(Layout!E2:E101),Layout!E2:E101,0)),"
      "MAX(1,n_min))", "modules", "The length that gets closest to the target.", "0",
      check=design.n_series)
    p(calc, "n_parallel", "Strings in parallel",
      "=IF(feasible,INDEX(Layout!B2:B101,MATCH(MIN(Layout!E2:E101),Layout!E2:E101,0)),1)",
      "strings", "", "0", check=design.n_parallel)
    p(calc, "n_modules", "Modules installed", "=n_series*n_parallel", "modules", "", "0",
      check=design.n_modules)
    p(calc, "dc_kwp", "Installed DC power", "=n_modules*p_stc_w/1000", "kWp", "", "0.000",
      check=design.dc_power_kwp)
    p(calc, "dc_ac", "DC/AC ratio", "=dc_kwp/p_ac_kw", "-", "", "0.00", check=design.dc_ac_ratio)
    p(calc, "voc_string", "Voc of a string, coldest", "=n_series*voc_cold", "V",
      "Must be below the inverter maximum voltage.", "0", check=design.voc_cold_string_v)
    p(calc, "vmp_string_hot", "Vmp of a string, hottest", "=n_series*vmp_hot", "V", "", "0",
      check=design.vmp_hot_string_v)
    p(calc, "vmp_string_cold", "Vmp of a string, coldest", "=n_series*vmp_cold", "V", "", "0",
      check=design.vmp_cold_string_v)
    book.heading(calc, "4. Yield")
    p(calc, "specific_yield", "Specific yield", "=irradiation*pr_effective", "kWh/kWp/year",
      "H x PR", "0.0", check=design.specific_yield_kwh_per_kwp)
    p(calc, "energy_kwh", "Yearly energy", "=dc_kwp*specific_yield", "kWh/year",
      "E = P x H x PR", "#,##0", check=design.annual_energy_kwh)
    p(calc, "achieved_coverage", "Share of the demand covered", "=energy_kwh/demand_kwh", "-", "",
      "0.0%", check=design.achieved_coverage)
    p(calc, "array_area", "Array area", "=n_modules*module_area", "m2", "", "0.0",
      check=design.array_area_m2)

    lay = book.sheet("Layout", (14, 14, 14, 14, 16))
    lay.append(["Modules in series", "Strings in parallel", "Modules", "Distance from target",
                "Score (lowest wins)"])
    for c in range(1, 6):
        lay.cell(1, c).font = _BOLD
        lay.cell(1, c).fill = _HEAD_FILL
    for n in range(1, _LAYOUT_ROWS + 1):
        r = n + 1
        lay.append([
            n,
            f"=MIN(string_capacity,MAX(1,ROUND(modules_target/A{r},0)))",
            f"=A{r}*B{r}",
            f"=ABS(C{r}-modules_target)",
            f"=IF(AND(A{r}>=MAX(1,n_min),A{r}<=n_max),D{r}*1000-A{r},1E+9)",
        ])
    lay.append([])
    lay.append(["Each row tries one string length. A length is allowed if it lies between the "
                "shortest and longest string; the best one has the module count closest to the "
                "target (ties go to the longer string)."])
    lay.freeze_panes = "A2"

    if inp.economics is not None and design.economics is not None:
        _economics(book, inp.economics, "dc_kwp", "energy_kwh", "demand_kwh", design.economics,
                   "kWp")
    book.finish_checks()
    return book.wb


# --- wind ----------------------------------------------------------------------------------
def build_wind_workbook(inp: WindInputs, design: WindDesign) -> Workbook:
    book = _Book()
    t = inp.turbine
    wake_note = ("The wake model is on in RESize: its result is used here as a fixed input "
                 "('array efficiency') and the number of turbines is the one RESize found."
                 if inp.wake_model else
                 "The wake model is off: the 'other losses' number stands for wakes and the rest.")
    _about(book, "RESize - wind design as a spreadsheet", _COMMON_ABOUT + [
        "Reproduced: wind speed at hub height, Weibull scale, air density, yield per turbine "
        "from the power curve (table on 'Energy'), number of turbines, economics.",
        wake_note,
    ])
    ws = book.sheet("Inputs")
    ws.append(["Inputs (yellow cells can be changed)"])
    ws.cell(1, 1).font = _BOLD
    p = book.put
    book.heading(ws, "Demand and wind resource")
    p(ws, "demand_kwh", "Annual electricity demand", inp.annual_demand_kwh, "kWh/year", "",
      "#,##0", True)
    p(ws, "coverage", "Share of the demand to cover", inp.coverage, "-", "", "0.00", True)
    p(ws, "v_ref", "Mean wind speed at the reference height", inp.mean_speed_ref_ms, "m/s", "",
      "0.00", True)
    p(ws, "h_ref", "Reference height", inp.ref_height_m, "m", "", "0.0", True)
    p(ws, "weibull_k", "Weibull shape factor k", inp.weibull_k, "-", "", "0.00", True)
    p(ws, "alpha", "Wind shear exponent", inp.shear_alpha, "-", "", "0.00", True)
    p(ws, "elevation", "Elevation of the site", inp.elevation_m, "m", "", "0", True)
    p(ws, "temperature", "Mean air temperature", inp.temperature_c, "degC", "", "0.0", True)
    p(ws, "availability", "Availability", inp.availability, "-", "", "0.00", True)
    p(ws, "other_losses", "Other losses", inp.other_losses, "-", "", "0.00", True)
    book.heading(ws, "Turbine (datasheet)")
    p(ws, "p_rated", "Rated power", t.rated_power_kw, "kW", t.name, "0", True)
    p(ws, "rotor_d", "Rotor diameter", t.rotor_diameter_m, "m", "", "0.0", True)
    p(ws, "hub_h", "Hub height", t.hub_height_m, "m", "", "0.0", True)
    p(ws, "cut_in", "Cut-in speed", t.cut_in_ms, "m/s", "", "0.0", True)
    p(ws, "v_rated", "Rated speed", t.rated_speed_ms, "m/s", "", "0.0", True)
    p(ws, "cut_out", "Cut-out speed", t.cut_out_ms, "m/s", "", "0.0", True)
    book.heading(ws, "From the wake model of RESize")
    eff = design.wake.array_efficiency if design.wake else 1.0
    p(ws, "array_eff", "Array efficiency (1 = no wake model)", eff, "-",
      "Not reproduced in the spreadsheet.", "0.000", True)
    if inp.wake_model:
        p(ws, "n_given", "Number of turbines found by RESize", design.n_turbines, "turbines",
          "With the wake model on, the number of turbines depends on the layout.", "0", True)

    calc = book.sheet("Calculation")
    calc.append(["Calculation (every cell is a formula)"])
    calc.cell(1, 1).font = _BOLD
    book.heading(calc, "1. Resource at hub height")
    p(calc, "v_hub", "Mean wind speed at hub height", "=v_ref*(hub_h/h_ref)^alpha", "m/s",
      "Power law: v = v_ref x (h / h_ref)^alpha", "0.000", check=design.mean_speed_hub_ms)
    p(calc, "weibull_a", "Weibull scale A", "=v_hub/EXP(GAMMALN(1+1/weibull_k))", "m/s",
      "A = mean speed / Gamma(1 + 1/k)", "0.000", check=design.weibull_scale_ms)
    p(calc, "rho", "Air density",
      "=101325*(1-2.25577E-5*elevation)^5.25588/(287.05*(temperature+273.15))", "kg/m3",
      "Pressure from elevation, ideal gas law", "0.0000", check=design.air_density)
    p(calc, "swept_area", "Swept area", "=PI()*rotor_d^2/4", "m2", "", "#,##0",
      check=design.swept_area_m2)
    p(calc, "rho_factor", "Density speed correction", "=(rho/1.225)^(1/3)", "-",
      "The power curve is read at v x (rho / 1.225)^(1/3)", "0.0000")
    book.heading(calc, "2. Yield of one turbine (see the 'Energy' sheet)")
    p(calc, "gross_kwh", "Gross yearly energy per turbine",
      "=8760*SUMPRODUCT(Energy!B2:B402,Energy!C2:C402)", "kWh/year",
      "E = 8760 x sum of P(v) x probability of that speed", "#,##0",
      check=design.gross_energy_per_turbine_kwh)
    p(calc, "net_free_kwh", "Net energy per turbine before wakes",
      "=gross_kwh*availability*(1-other_losses)", "kWh/year", "", "#,##0")
    p(calc, "net_kwh", "Net yearly energy per turbine", "=net_free_kwh*array_eff", "kWh/year",
      "", "#,##0", check=design.net_energy_per_turbine_kwh)
    p(calc, "cf_net", "Capacity factor (net)", "=net_kwh/(p_rated*8760)", "-", "", "0.0%",
      check=design.capacity_factor_net)
    book.heading(calc, "3. How many turbines")
    count = ("=n_given" if inp.wake_model else
             "=MAX(1,ROUNDUP(coverage*demand_kwh/net_free_kwh-1E-9,0))")
    p(calc, "n_turbines", "Number of turbines", count, "turbines",
      "Demand to cover / energy per turbine, rounded up", "0", check=design.n_turbines)
    p(calc, "power_kw", "Installed power", "=n_turbines*p_rated", "kW", "", "#,##0",
      check=design.installed_power_kw)
    p(calc, "energy_kwh", "Yearly energy of the plant", "=n_turbines*net_kwh", "kWh/year", "",
      "#,##0", check=design.annual_energy_kwh)
    p(calc, "achieved_coverage", "Share of the demand covered", "=energy_kwh/demand_kwh", "-", "",
      "0.0%", check=design.achieved_coverage)
    p(calc, "specific_power", "Specific power", "=p_rated*1000/swept_area", "W/m2", "", "0.0",
      check=design.specific_power_w_m2)

    en = book.sheet("Energy", (14, 16, 18, 18))
    en.append(["Wind speed (m/s)", "Power (kW)", "Probability of the bin"])
    for c in range(1, 5):
        en.cell(1, c).font = _BOLD
        en.cell(1, c).fill = _HEAD_FILL
    curve = t.power_curve
    if curve:
        tab = book.wb.create_sheet("Power curve")
        tab.append(["Speed (m/s)", "Power (kW)"])
        for c in (1, 2):
            tab.cell(1, c).font = _BOLD
        for vs, ps in curve:
            tab.append([vs, ps])
        last = len(curve) + 1
    for i in range(401):
        r = i + 2
        v = f"A{r}"
        veq = f"({v}*rho_factor)"
        if curve:
            power = (f"=IF(OR({v}>=cut_out,{veq}<'Power curve'!$A$2),0,"
                     f"IF({veq}>='Power curve'!$A${last},'Power curve'!$B${last},"
                     f"INDEX('Power curve'!$B$2:$B${last},MATCH({veq},'Power curve'!$A$2:$A${last},1))"
                     f"+(INDEX('Power curve'!$B$2:$B${last},MATCH({veq},'Power curve'!$A$2:$A${last},1)+1)"
                     f"-INDEX('Power curve'!$B$2:$B${last},MATCH({veq},'Power curve'!$A$2:$A${last},1)))"
                     f"*({veq}-INDEX('Power curve'!$A$2:$A${last},MATCH({veq},'Power curve'!$A$2:$A${last},1)))"
                     f"/(INDEX('Power curve'!$A$2:$A${last},MATCH({veq},'Power curve'!$A$2:$A${last},1)+1)"
                     f"-INDEX('Power curve'!$A$2:$A${last},MATCH({veq},'Power curve'!$A$2:$A${last},1)))))")
        else:
            power = (f"=IF(OR({v}>=cut_out,{veq}<=cut_in),0,IF({veq}>=v_rated,p_rated,"
                     f"p_rated*({veq}^3-cut_in^3)/(v_rated^3-cut_in^3)))")
        top = f"(1-EXP(-((({v}+0.05)/weibull_a)^weibull_k)))"
        bottom = f"IF({v}-0.05<=0,0,1-EXP(-((({v}-0.05)/weibull_a)^weibull_k)))"
        en.append([round(i * 0.1, 1), power, f"={top}-{bottom}", None])
        en.cell(r, 3).number_format = "0.000000"
        en.cell(r, 2).number_format = "0.0"
    en.freeze_panes = "A2"

    if inp.economics is not None and design.economics is not None:
        _economics(book, inp.economics, "power_kw", "energy_kwh", "demand_kwh", design.economics,
                   "kW")
    book.finish_checks()
    return book.wb


def save_workbook(wb: Workbook, path: str | Path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))

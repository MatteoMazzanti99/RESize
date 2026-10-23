"""Excel workbooks with live formulas for biomass, hydrogen, battery and solar thermal.

Qt-free; builds on the helpers of :mod:`resize_energy.report.excel`.

* Biomass: the whole calculation is formulas (the load-duration coverage is solved by a
  bisection table, as in the program).
* Battery and solar thermal: the hour-by-hour model is rebuilt row by row (8760 rows of
  formulas). Load, production and weather columns are pasted as values: they come from
  generators with random numbers that a spreadsheet cannot repeat.
* Hydrogen: the hourly dispatch is rebuilt too (renewable power, grid top-up, minimum load,
  part-load efficiency, tank swing, ageing); everything downstream is formulas.
"""

from __future__ import annotations

from openpyxl import Workbook

from resize_energy.core import battery as bat
from resize_energy.core import hydrogen as h2
from resize_energy.core import solar_thermal as st
from resize_energy.core.battery_design import BatteryDesign, BatteryInputs, BatteryProfiles
from resize_energy.core.battery_design import _generation as battery_generation
from resize_energy.core.biomass_design import BiomassDesign, BiomassInputs
from resize_energy.core.hydrogen_design import HydrogenDesign, HydrogenInputs, HydrogenProfiles
from resize_energy.core.solar_thermal_design import SolarThermalDesign, SolarThermalInputs
from resize_energy.report.excel import _BOLD, _COMMON_ABOUT, _HEAD_FILL, _about, _Book

_HOURS = 8760


def _crf(rate: str, years: str) -> str:
    return (f"IF({rate}=0,1/({years}),{rate}*(1+{rate})^({years})/((1+{rate})^({years})-1))")


def _annuity(rate: str, years: str) -> str:
    return f"(1/{_crf(rate, years)})"


def _table_header(ws, names: list[str]) -> None:
    ws.append(names)
    for c in range(1, len(names) + 1):
        ws.cell(1, c).font = _BOLD
        ws.cell(1, c).fill = _HEAD_FILL


def _title(ws, text: str) -> None:
    ws.append([text])
    ws.cell(1, 1).font = _BOLD


# --- biomass -------------------------------------------------------------------------------
def build_biomass_workbook(inp: BiomassInputs, d: BiomassDesign) -> Workbook:
    book = _Book()
    p = book.put
    solid = inp.pathway != "biogas"
    _about(book, "RESize - biomass design as a spreadsheet", _COMMON_ABOUT + [
        "Everything is reproduced, including the share of the heat covered by a plant of a given "
        "size (sheet 'Coverage': a bisection, one row per step)." if solid else
        "Everything is reproduced: methane yield of the substrate mix, digester volume, "
        "engine, heat balance, land and emissions.",
    ])
    ws = book.sheet("Inputs")
    _title(ws, "Inputs (yellow cells can be changed)")
    book.heading(ws, "Demand")
    if solid:
        p(ws, "q_heat", "Annual heat demand", inp.annual_heat_demand_kwh, "kWh/year", "",
          "#,##0", True)
        p(ws, "peak_kw", "Peak heat load", inp.peak_heat_kw, "kW", "", "#,##0.0", True)
        p(ws, "el_demand", "Annual electricity demand", inp.annual_electricity_demand_kwh,
          "kWh/year", "Only used to compare with the electricity of a cogeneration plant.",
          "#,##0", True)
    else:
        p(ws, "el_demand", "Annual electricity demand", inp.annual_electricity_demand_kwh,
          "kWh/year", "The biogas plant is sized on the electricity.", "#,##0", True)
        p(ws, "q_heat", "Annual heat demand", inp.annual_heat_demand_kwh, "kWh/year",
          "To see how much of the CHP heat finds a use.", "#,##0", True)
    p(ws, "coverage", "Share of the driving demand to cover", inp.coverage, "-", "", "0.00", True)
    p(ws, "eta_th", "Thermal efficiency", inp.eta_th, "-", "Useful heat / fuel energy",
      "0.000", True)
    p(ws, "eta_el", "Electrical efficiency", inp.eta_el, "-", "", "0.000", True)
    if solid:
        f = inp.fuel
        book.heading(ws, "Fuel")
        p(ws, "lhv_dry", "LHV of the dry matter", f.lhv_dry_kwh_per_kg, "kWh/kg", f.name,
          "0.000", True)
        p(ws, "moisture", "Moisture, wet basis", f.moisture_wb, "-", "", "0.00", True)
        p(ws, "bulk", "Bulk density as delivered", f.bulk_density_kg_m3, "kg/m3", "", "0", True)
        p(ws, "ash_dry", "Ash, share of dry matter", f.ash_dry, "-", "", "0.000", True)
        p(ws, "storage_days", "Fuel storage autonomy", inp.storage_days, "days", "", "0.0", True)
        p(ws, "truck_m3", "Truck volume", inp.truck_volume_m3, "m3", "", "0.0", True)
    else:
        a, b = inp.substrate_a, inp.substrate_b
        book.heading(ws, "Substrates")
        for tag, s in (("a", a), ("b", b)):
            p(ws, f"ch4_{tag}", f"Methane per tonne, substrate {tag.upper()}", s.ch4_m3_per_t,
              "Nm3/t", s.name, "0.0", True)
            p(ws, f"frac_{tag}", f"CH4 share of the biogas, {tag.upper()}", s.ch4_fraction, "-",
              "", "0.00", True)
            p(ws, f"vs_{tag}", f"Volatile solids (fresh matter), {tag.upper()}", s.vs_fresh,
              "-", "TS x VS/TS", "0.000", True)
            p(ws, f"yield_{tag}", f"Crop yield, {tag.upper()}", s.crop_yield_t_per_ha, "t/ha",
              "0 for wastes and manure", "0.0", True)
        p(ws, "share_b", "Share of substrate B in the feed (mass)", inp.share_b, "-", "", "0.00",
          True)
        book.heading(ws, "Digester and engine")
        p(ws, "hrt", "Hydraulic retention time", inp.hrt_days, "days", "", "0", True)
        p(ws, "op_hours", "CHP full-load hours", inp.operating_hours, "h/year", "", "0", True)
        p(ws, "parasitic", "Plant's own electricity use", inp.parasitic_electricity, "-", "",
          "0.00", True)
        p(ws, "dig_share", "CHP heat used by the digester", inp.digester_heat_share, "-", "",
          "0.00", True)
    book.heading(ws, "Emissions")
    p(ws, "ref_eff", "Efficiency of the heat it replaces", inp.reference_heat_efficiency, "-", "",
      "0.00", True)
    p(ws, "ref_ef", "Emission factor of that heat", inp.reference_heat_ef, "kg CO2/kWh", "",
      "0.000", True)
    p(ws, "grid_ef", "Emission factor of the grid", inp.grid_ef, "kg CO2/kWh", "", "0.000", True)
    p(ws, "supply_ef", "Supply-chain emissions", inp.supply_chain_ef, "kg CO2/kWh fuel",
      "Biogenic CO2 of the fuel counts as neutral.", "0.000", True)

    c = book.sheet("Calculation")
    _title(c, "Calculation (every cell is a formula)")
    if solid:
        book.heading(c, "1. Size of the plant")
        p(c, "load_factor", "Load factor of the heat demand", "=q_heat/(peak_kw*8760)", "-",
          "Mean load / peak load", "0.0000", check=d.load_factor)
        p(c, "n_exp", "Exponent of the load-duration curve", "=1/load_factor-1", "-",
          "P(t) = P_peak x (1 - t/T)^n", "0.0000")
        p(c, "power_ratio", "Plant power / peak load",
          "=IF(coverage>=1,1,(Coverage!A82+Coverage!B82)/2)", "-",
          "Found on the 'Coverage' sheet: the ratio that covers the requested share.", "0.0000",
          check=d.power_ratio)
        p(c, "p_th", "Thermal power of the plant", "=power_ratio*peak_kw", "kW", "", "#,##0.0",
          check=d.thermal_power_kw)
        p(c, "heat", "Heat delivered", "=coverage*q_heat", "kWh/year", "", "#,##0",
          check=d.heat_useful_kwh)
        p(c, "fuel_kwh", "Fuel energy", "=heat/eta_th", "kWh/year", "Heat / efficiency",
          "#,##0", check=d.fuel_energy_kwh)
        p(c, "el_kwh", "Electricity (cogeneration)", "=fuel_kwh*eta_el", "kWh/year", "",
          "#,##0", check=d.electricity_net_kwh)
        p(c, "el_power", "Electric power", "=p_th*eta_el/eta_th", "kW", "", "#,##0.0",
          check=d.electric_power_kw)
        p(c, "flh", "Full-load hours", "=heat/p_th", "h/year", "", "#,##0",
          check=d.full_load_hours)
        book.heading(c, "2. Fuel")
        p(c, "lhv", "LHV as delivered",
          f"=lhv_dry*(1-moisture)-{2.443 / 3.6!r}*moisture", "kWh/kg",
          "LHV_dry (1 - M) - 0.679 M", "0.0000", check=d.lhv_as_received_kwh_per_kg)
        p(c, "mass_t", "Fuel mass", "=fuel_kwh/lhv/1000", "t/year", "", "#,##0.0",
          check=d.mass_t_per_year)
        p(c, "ash_t", "Ash", "=mass_t*(1-moisture)*ash_dry", "t/year", "", "#,##0.00",
          check=d.ash_t_per_year)
        p(c, "storage_m3", "Fuel store",
          "=p_th/eta_th/lhv*24*storage_days/bulk/0.8", "m3",
          "Consumption at nominal power for the autonomy, bin 80 % usable", "#,##0.0",
          check=d.storage_m3)
        p(c, "deliveries", "Truck deliveries", "=ROUNDUP(mass_t/(truck_m3*bulk/1000),0)",
          "per year", "", "0", check=d.deliveries_per_year)
        book.heading(c, "3. Emissions")
        p(c, "co2_t_year", "CO2 avoided",
          "=(heat/ref_eff*ref_ef+el_kwh*grid_ef-fuel_kwh*supply_ef)/1000", "t CO2/year",
          "Replaced emissions minus supply chain", "#,##0.0", check=d.co2_avoided_t_per_year)
        cov = book.sheet("Coverage", (14, 14, 14, 16))
        _table_header(cov, ["Low", "High", "Trial ratio", "Heat covered by the trial ratio"])
        for i in range(81):
            r = i + 2
            if i == 0:
                lo, hi = 0, 1
            else:
                lo = f"=IF(D{r - 1}<coverage,C{r - 1},A{r - 1})"
                hi = f"=IF(D{r - 1}<coverage,B{r - 1},C{r - 1})"
            cov.append([lo, hi, f"=(A{r}+B{r})/2",
                        f"=IF(n_exp<1E-9,C{r},(n_exp+1)*C{r}-n_exp*C{r}^((n_exp+1)/n_exp))"])
        cov.append([])
        cov.append(["Bisection: the covered share f(x) = (n+1)x - n x^((n+1)/n) rises with the "
                    "plant size x; each row halves the interval that contains the answer."])
    else:
        book.heading(c, "1. Energy chain")
        p(c, "el_net", "Net electricity to deliver", "=coverage*el_demand", "kWh/year", "",
          "#,##0", check=d.electricity_net_kwh)
        p(c, "el_gross", "Gross electricity", "=el_net/(1-parasitic)", "kWh/year",
          "Plus the plant's own use", "#,##0", check=d.electricity_gross_kwh)
        p(c, "fuel_kwh", "Methane energy", "=el_gross/eta_el", "kWh/year", "", "#,##0",
          check=d.fuel_energy_kwh)
        p(c, "ch4_m3", "Methane", f"=fuel_kwh/{bm_ch4()}", "Nm3/year", "LHV 9.97 kWh/Nm3",
          "#,##0", check=d.ch4_m3_per_year)
        book.heading(c, "2. Substrate")
        p(c, "ch4_per_t", "Methane per tonne of the mix", "=(1-share_b)*ch4_a+share_b*ch4_b",
          "Nm3/t", "", "0.00")
        p(c, "biogas_per_t", "Biogas per tonne of the mix",
          "=(1-share_b)*ch4_a/frac_a+share_b*ch4_b/frac_b", "Nm3/t", "", "0.00")
        p(c, "mass_t", "Substrate", "=ch4_m3/ch4_per_t", "t/year", "", "#,##0.0",
          check=d.mass_t_per_year)
        p(c, "mass_a", "  of which A", "=(1-share_b)*mass_t", "t/year", "", "#,##0.0",
          check=d.substrate_mass_a_t)
        p(c, "mass_b", "  of which B", "=share_b*mass_t", "t/year", "", "#,##0.0",
          check=d.substrate_mass_b_t)
        p(c, "biogas_m3", "Biogas", "=mass_t*biogas_per_t", "Nm3/year", "", "#,##0",
          check=d.biogas_m3_per_year)
        p(c, "ch4_frac", "CH4 share of the biogas", "=ch4_m3/biogas_m3", "-", "", "0.000",
          check=d.ch4_fraction)
        book.heading(c, "3. Digester")
        p(c, "per_day", "Feed per day", "=mass_t/365", "t/day", "", "0.00")
        p(c, "v_work", "Working volume", "=per_day*hrt", "m3", "Feed x retention time", "#,##0",
          check=d.working_volume_m3)
        p(c, "v_total", "Digester volume", "=per_day*hrt/0.9", "m3", "Fill fraction 0.9",
          "#,##0", check=d.digester_volume_m3)
        p(c, "olr", "Organic loading rate",
          "=(mass_a*vs_a+mass_b*vs_b)*1000/365/v_work", "kg VS/m3/day", "", "0.00",
          check=d.olr_kg_vs_per_m3_day)
        p(c, "hectares", "Land for the crops",
          "=IF(yield_a>0,mass_a/yield_a,0)+IF(yield_b>0,mass_b/yield_b,0)", "ha", "", "#,##0.0",
          check=d.crop_hectares)
        p(c, "digestate", "Digestate", f"=mass_t-biogas_m3*{bm_density()}/1000", "t/year", "",
          "#,##0.0", check=d.digestate_t_per_year)
        book.heading(c, "4. Heat and emissions")
        p(c, "heat_total", "CHP heat", "=el_gross*eta_th/eta_el", "kWh/year", "", "#,##0",
          check=d.heat_total_kwh)
        p(c, "heat_dig", "Heat for the digester", "=dig_share*heat_total", "kWh/year", "",
          "#,##0", check=d.heat_digester_kwh)
        p(c, "heat_use", "Heat available", "=heat_total-heat_dig", "kWh/year", "", "#,##0",
          check=d.heat_useful_kwh)
        p(c, "credited", "Heat that replaces fossil heat",
          "=IF(q_heat>0,MIN(heat_use,q_heat),0)", "kWh/year", "", "#,##0",
          check=d.heat_credited_kwh)
        p(c, "co2_t_year", "CO2 avoided",
          "=(credited/ref_eff*ref_ef+el_net*grid_ef-fuel_kwh*supply_ef)/1000", "t CO2/year", "",
          "#,##0.0", check=d.co2_avoided_t_per_year)
        p(c, "el_power", "Electric power", "=el_gross/op_hours", "kW", "", "#,##0.0",
          check=d.electric_power_kw)
    book.finish_checks()
    return book.wb


def bm_ch4() -> str:
    from resize_energy.core import biomass as bm
    return repr(bm.CH4_LHV_KWH_PER_NM3)


def bm_density() -> str:
    from resize_energy.core import biomass as bm
    return repr(bm.BIOGAS_DENSITY_KG_PER_NM3)


# --- battery -------------------------------------------------------------------------------
def build_battery_workbook(inp: BatteryInputs, prof: BatteryProfiles, d: BatteryDesign
                           ) -> Workbook:
    book = _Book()
    p = book.put
    load = bat.load_profile(inp.annual_demand_kwh, inp.load_kind)
    gen = battery_generation(inp, prof)
    _about(book, "RESize - battery as a spreadsheet", _COMMON_ABOUT + [
        "The sheet 'Hourly' rebuilds the 8760-hour model with formulas: serve the load, charge "
        "from the surplus, discharge into the deficit, import and export.",
        "The load and production columns are values (they come from the program's profile "
        "generators, which use random numbers); change them if you have measured data.",
        "The capacity is the one chosen by RESize. Type another capacity to try it: the target "
        "and optimum searches of the program are not repeated.",
    ])
    ws = book.sheet("Inputs")
    _title(ws, "Inputs (yellow cells can be changed)")
    book.heading(ws, "Battery")
    p(ws, "capacity", "Capacity", d.capacity_kwh, "kWh", "", "0.00", True)
    p(ws, "c_rate", "C-rate (kW per kWh)", inp.c_rate, "kW/kWh", "", "0.00", True)
    p(ws, "eta_rt", "Round-trip efficiency", inp.eta_roundtrip, "-", "", "0.00", True)
    p(ws, "dod", "Usable fraction (depth of discharge)", inp.usable_fraction, "-", "", "0.00",
      True)
    p(ws, "health", "Average state of health", inp.health, "-", "", "0.00", True)
    p(ws, "soc0", "Initial charge", 0.5, "-", "Share of the usable energy", "0.00", True)
    p(ws, "cycle_life", "Cycle life", inp.cycle_life, "cycles", "", "0", True)
    p(ws, "calendar_life", "Calendar life", inp.calendar_life_years, "years", "", "0.0", True)
    book.heading(ws, "Costs and prices")
    p(ws, "capex_kwh", "Cost per kWh", inp.capex_per_kwh, "euro/kWh", "", "0.00", True)
    p(ws, "capex_kw", "Cost per kW", inp.capex_per_kw, "euro/kW", "", "0.00", True)
    p(ws, "om_fraction", "O&M, share of the cost per year", inp.om_fraction, "-", "", "0.000",
      True)
    p(ws, "retail", "Price you pay for electricity", inp.retail_price, "euro/kWh", "", "0.000",
      True)
    p(ws, "export_price", "Price paid for the surplus", inp.export_price, "euro/kWh", "",
      "0.000", True)
    p(ws, "disc_rate", "Discount rate", inp.discount_rate, "-", "", "0.0%", True)
    p(ws, "grid_ef", "Emission factor of the grid", inp.grid_ef, "kg CO2/kWh", "", "0.000", True)

    c = book.sheet("Calculation")
    _title(c, "Calculation (every cell is a formula)")
    book.heading(c, "1. The battery")
    p(c, "usable", "Usable capacity", "=capacity*dod", "kWh", "", "0.00",
      check=d.usable_kwh)
    p(c, "power_kw", "Power", "=capacity*c_rate", "kW", "", "0.00", check=d.power_kw)
    p(c, "room", "Energy it can really hold", "=capacity*dod*health", "kWh",
      "Capacity x usable fraction x health", "0.00")
    p(c, "eta", "Efficiency each way", "=SQRT(eta_rt)", "-",
      "The same loss when charging and discharging", "0.0000")
    p(c, "soc_start", "Initial energy", "=room*soc0", "kWh", "", "0.00")
    book.heading(c, "2. The year (sums of the 'Hourly' sheet)")
    n = _HOURS + 1
    p(c, "load_kwh", "Load", f"=SUM(Hourly!B2:B{n})", "kWh", "", "#,##0", check=d.load_kwh)
    p(c, "gen_kwh", "Production", f"=SUM(Hourly!C2:C{n})", "kWh", "", "#,##0",
      check=d.generation_kwh)
    p(c, "direct", "Used directly", f"=SUM(Hourly!D2:D{n})", "kWh", "", "#,##0",
      check=d.direct_kwh)
    p(c, "charge", "Sent to the battery", f"=SUM(Hourly!F2:F{n})", "kWh", "", "#,##0")
    p(c, "discharge", "Delivered by the battery", f"=SUM(Hourly!G2:G{n})", "kWh", "", "#,##0",
      check=d.battery_kwh)
    p(c, "import_kwh", "Imported from the grid", f"=SUM(Hourly!I2:I{n})", "kWh", "", "#,##0",
      check=d.import_kwh)
    p(c, "export_kwh", "Exported to the grid", f"=SUM(Hourly!J2:J{n})", "kWh", "", "#,##0",
      check=d.export_kwh)
    p(c, "import0", "Imported without the battery", f"=SUM(Hourly!K2:K{n})", "kWh", "",
      "#,##0", check=d.import_without_kwh)
    p(c, "export0", "Exported without the battery", f"=SUM(Hourly!L2:L{n})", "kWh", "",
      "#,##0", check=d.export_without_kwh)
    p(c, "loss", "Lost in the battery", "=charge-discharge", "kWh", "Round-trip loss", "#,##0",
      check=d.loss_kwh)
    book.heading(c, "3. Indicators")
    p(c, "self_suff", "Self-sufficiency", "=1-import_kwh/load_kwh", "-",
      "Share of the load met without the grid", "0.0%", check=d.self_sufficiency)
    p(c, "self_suff0", "Self-sufficiency without battery", "=1-import0/load_kwh", "-", "",
      "0.0%", check=d.self_sufficiency_without)
    p(c, "self_cons", "Self-consumption", "=(direct+charge)/gen_kwh", "-",
      "Share of the production used on site", "0.0%", check=d.self_consumption)
    p(c, "self_cons0", "Self-consumption without battery", "=direct/gen_kwh", "-", "", "0.0%",
      check=d.self_consumption_without)
    book.heading(c, "4. Life and economics")
    p(c, "cycles", "Equivalent full cycles per year", "=IF(usable>0,discharge/usable,0)",
      "cycles/year", "", "0.0", check=d.cycles_per_year)
    p(c, "life_cycles", "Life from cycles", "=IF(cycles>0,cycle_life/cycles,1E+99)", "years",
      "", "0.0")
    p(c, "life", "Life of the battery", "=MIN(calendar_life,life_cycles)", "years",
      "The sooner of cycle life and calendar life", "0.0", check=d.life_years)
    p(c, "capex", "Cost (CAPEX)", "=capacity*capex_kwh+power_kw*capex_kw", "euro", "", "#,##0",
      check=d.capex)
    p(c, "om", "O&M per year", "=capex*om_fraction", "euro/year", "", "#,##0",
      check=d.om_per_year)
    p(c, "savings", "Savings per year",
      "=(import0-import_kwh)*retail-(export0-export_kwh)*export_price", "euro/year",
      "Less import at the retail price, less export at the export price", "#,##0",
      check=d.savings_per_year)
    p(c, "net", "Net benefit per year", "=savings-om", "euro/year", "", "#,##0",
      check=d.net_benefit_per_year)
    p(c, "payback", "Simple payback", '=IF(net>0,capex/net,"never")', "years", "", "0.0",
      check=d.payback_years)
    p(c, "npv_euro", "Net present value",
      f"=IF(capacity>0,-capex+net*{_annuity('disc_rate', 'life')},0)", "euro",
      "-CAPEX + net benefit x annuity factor over the battery life", "#,##0", check=d.npv)
    p(c, "lcos", "Cost of the energy delivered",
      f"=IF(AND(discharge>0,capacity>0),(capex*{_crf('disc_rate', 'life')}+om)/discharge,0)",
      "euro/kWh", "(CAPEX x CRF + O&M) / energy delivered", "0.000", check=d.lcos)
    p(c, "co2_t_year", "CO2 avoided", "=(import0-import_kwh)*grid_ef/1000", "t CO2/year", "",
      "#,##0.00", check=d.co2_avoided_t_per_year)

    h = book.sheet("Hourly", (8, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12))
    _table_header(h, ["Hour", "Load (kW)", "Production (kW)", "Direct use", "Surplus (+) or "
                      "deficit (-)", "To battery", "From battery", "Energy stored", "Import",
                      "Export", "Import (no battery)", "Export (no battery)"])
    for t in range(_HOURS):
        r = t + 2
        prev = "soc_start" if t == 0 else f"H{r - 1}"
        h.append([
            t, float(load[t]), float(gen[t]), f"=MIN(B{r},C{r})", f"=C{r}-B{r}",
            f"=IF(E{r}>0,IF(room>0,MAX(0,MIN(E{r},power_kw,(room-{prev})/eta)),0),0)",
            f"=IF(E{r}<0,IF(room>0,MAX(0,MIN(-E{r},power_kw,{prev}*eta)),0),0)",
            f"={prev}+F{r}*eta-G{r}/eta",
            f"=IF(E{r}<0,-E{r}-G{r},0)", f"=IF(E{r}>0,E{r}-F{r},0)",
            f"=MAX(-E{r},0)", f"=MAX(E{r},0)"])
    h.freeze_panes = "A2"
    book.finish_checks()
    return book.wb


# --- solar thermal -------------------------------------------------------------------------
def build_solar_workbook(inp: SolarThermalInputs, d: SolarThermalDesign) -> Workbook:
    book = _Book()
    p = book.put
    s_inp = d.inputs_used or inp
    weather = st.synthetic_weather(s_inp.latitude_deg, s_inp.irradiation_kwh_m2, s_inp.t_mean_c,
                                   s_inp.t_season_amp_c)
    col = d.collector
    litres = d.store_litres
    run = d.run
    _about(book, "RESize - solar thermal as a spreadsheet", _COMMON_ABOUT + [
        "The sheet 'Hourly' rebuilds the hour-by-hour model with formulas: collector loop, "
        "stagnation limit, hot-water draw, heat-driven chiller and standing loss of the store.",
        "Irradiance and air temperature are values: the program invents a physical year with a "
        "random generator that a spreadsheet cannot repeat. Paste measured data to replace them.",
        "The collector area and the chiller size are the ones found by RESize. Only the main "
        "cooling technology is rebuilt; the comparison of technologies is not repeated here.",
    ])
    ws = book.sheet("Inputs")
    _title(ws, "Inputs (yellow cells can be changed)")
    book.heading(ws, "Site")
    p(ws, "t_mean", "Mean air temperature", s_inp.t_mean_c, "degC", "", "0.0", True)
    p(ws, "t_amp", "Seasonal amplitude", s_inp.t_season_amp_c, "degC", "", "0.0", True)
    book.heading(ws, "Collectors and store")
    p(ws, "area", "Collector area", d.area_m2, "m2", col.name, "0.0", True)
    p(ws, "coll_eta0", "Zero-loss efficiency", col.eta0, "-", "", "0.000", True)
    p(ws, "loss_k1", "Heat-loss coefficient a1", col.a1, "W/m2/K", "", "0.000", True)
    p(ws, "loss_k2", "Heat-loss coefficient a2", col.a2, "W/m2/K2", "", "0.0000", True)
    p(ws, "loop_loss", "Pipe and pump losses", s_inp.loop_loss, "-", "", "0.00", True)
    p(ws, "flow_dt", "Collector outlet minus store temperature", 6.0, "K", "", "0.0", True)
    p(ws, "store_l", "Store volume", litres, "L", "", "0", True)
    p(ws, "t_room", "Room temperature around the store", 18.0, "degC", "", "0.0", True)
    book.heading(ws, "Loads")
    p(ws, "dhw_l", "Hot water per day", s_inp.dhw_litres_per_day, "L/day", "", "0", True)
    p(ws, "t_use", "Hot-water temperature", s_inp.t_use_c, "degC", "", "0.0", True)
    p(ws, "cool_kwh", "Cooling demand", s_inp.cooling_kwh_per_year, "kWh/year", "", "#,##0",
      True)
    book.heading(ws, "Heat-driven chiller")
    p(ws, "chiller_kw", "Cooling power", d.chiller_kw, "kW", "", "0.0", True)
    p(ws, "cop_nom", "COP at the nominal drive temperature", s_inp.cop_nominal, "-", "",
      "0.00", True)
    p(ws, "t_min", "Minimum drive temperature", s_inp.t_drive_min_c, "degC", "", "0.0", True)
    p(ws, "t_nom", "Nominal drive temperature", s_inp.t_drive_nominal_c, "degC", "", "0.0",
      True)
    p(ws, "parasitic", "Electricity per kWh of cold", s_inp.chiller_parasitic, "kWh/kWh", "",
      "0.000", True)
    p(ws, "comp_cop", "COP of the conventional chiller it replaces", s_inp.compressor_cop, "-",
      "", "0.00", True)
    book.heading(ws, "Costs and prices")
    p(ws, "col_cost", "Collector cost", col.cost_per_m2, "euro/m2", "", "0.00", True)
    p(ws, "store_cost", "Store cost", s_inp.store_cost_per_litre, "euro/L", "", "0.00", True)
    p(ws, "chiller_cost", "Chiller cost", s_inp.chiller_cost_per_kw, "euro/kW", "", "0.00",
      True)
    p(ws, "fixed_cost", "Fixed cost", s_inp.fixed_cost, "euro", "", "#,##0", True)
    p(ws, "om_fraction", "O&M, share of the cost per year", s_inp.om_fraction, "-", "", "0.000",
      True)
    p(ws, "heat_price", "Price of backup heat", s_inp.heat_price, "euro/kWh", "", "0.000", True)
    p(ws, "elec_price", "Price of electricity", s_inp.electricity_price, "euro/kWh", "",
      "0.000", True)
    p(ws, "life_years", "Lifetime", s_inp.life_years, "years", "", "0.0", True)
    p(ws, "disc_rate", "Discount rate", s_inp.discount_rate, "-", "", "0.0%", True)
    p(ws, "heat_ef", "Emission factor of backup heat", s_inp.heat_ef_kg_per_kwh, "kg CO2/kWh",
      "", "0.000", True)
    p(ws, "grid_ef", "Emission factor of the grid", s_inp.grid_ef_kg_per_kwh, "kg CO2/kWh", "",
      "0.000", True)

    c = book.sheet("Calculation")
    _title(c, "Calculation (every cell is a formula)")
    n = _HOURS + 1
    book.heading(c, "1. The store")
    p(c, "cap", "Heat capacity of the store", f"=store_l*{st.WATER_KWH_PER_L_K}", "kWh/K",
      "Volume x 0.001163", "0.0000", check=run.store_capacity_kwh_k)
    p(c, "ua", "Standing-loss coefficient", "=0.16*SQRT(store_l)", "W/K", "UA = 0.16 sqrt(V)",
      "0.000")
    p(c, "heat_cap", "Heat the chiller can take", "=IF(cop_nom>0,chiller_kw/cop_nom,0)", "kW",
      "", "0.00")
    p(c, "raw_total", "Sum of the cooling shape", f"=SUM(Hourly!G2:G{n})", "K h",
      "Hours 8-20 above 18 degC, summed", "0.0")
    p(c, "t_start", "Initial store temperature", "=MAX(20,t_mean)", "degC", "", "0.0",
      check=run.t_start_c)
    book.heading(c, "2. The year (sums of the 'Hourly' sheet)")
    p(c, "collector_kwh", "Heat from the collectors", f"=SUM(Hourly!J2:J{n})-SUM(Hourly!L2:L{n})",
      "kWh", "After the stagnation dump", "#,##0", check=run.collector_kwh)
    p(c, "wasted", "Heat dumped at the stagnation limit", f"=SUM(Hourly!L2:L{n})", "kWh", "",
      "#,##0", check=run.wasted_kwh)
    p(c, "heat_dem", "Hot-water demand", f"=SUM(Hourly!F2:F{n})", "kWh", "", "#,##0",
      check=run.heat_demand_kwh)
    p(c, "heat_sol", "Hot water covered by the sun", f"=SUM(Hourly!P2:P{n})", "kWh", "",
      "#,##0", check=run.heat_solar_kwh)
    p(c, "cool_dem", "Cooling demand", f"=SUM(Hourly!H2:H{n})", "kWh", "", "#,##0",
      check=run.cooling_demand_kwh)
    p(c, "cool_sol", "Cold from the sun", f"=SUM(Hourly!T2:T{n})", "kWh", "", "#,##0",
      check=run.cooling_solar_kwh)
    p(c, "chiller_heat", "Heat used by the chiller", f"=SUM(Hourly!S2:S{n})", "kWh", "",
      "#,##0", check=run.chiller_heat_kwh)
    p(c, "store_loss", "Standing loss of the store", f"=SUM(Hourly!W2:W{n})", "kWh", "",
      "#,##0", check=run.store_loss_kwh)
    p(c, "stag_hours", "Hours at the stagnation limit", f"=SUM(Hourly!N2:N{n})", "h", "", "0",
      check=run.hours_stagnation)
    p(c, "below_hours", "Hours the store was too cold for the chiller",
      f"=SUM(Hourly!U2:U{n})", "h", "", "0", check=run.hours_below_drive)
    p(c, "t_end", "Store temperature at the end of the year", f"=Hourly!X{n}", "degC", "",
      "0.0", check=run.t_end_c)
    p(c, "heat_fraction", "Solar fraction of the hot water",
      '=IF(heat_dem>0,heat_sol/heat_dem,0)', "-", "", "0.0%", check=run.heat_fraction)
    p(c, "solar_fraction", "Solar fraction (heat and cold)",
      "=IF(heat_dem+cool_dem/cop_nom>0,(heat_sol+cool_sol/cop_nom)/(heat_dem+cool_dem/cop_nom),0)",
      "-", "Cold counted as the heat a boiler-driven chiller would need",
      "0.0%", check=run.solar_fraction)
    book.heading(c, "3. Costs, savings and carbon")
    p(c, "collector_cost", "Collectors", "=area*col_cost", "euro", "", "#,##0",
      check=d.collector_cost)
    p(c, "store_cost_total", "Store", "=store_l*store_cost", "euro", "", "#,##0",
      check=d.store_cost)
    p(c, "chiller_cost_total", "Chiller", "=chiller_kw*chiller_cost", "euro", "", "#,##0",
      check=d.chiller_cost)
    p(c, "capex", "Total cost (CAPEX)",
      "=collector_cost+store_cost_total+chiller_cost_total+fixed_cost", "euro", "", "#,##0",
      check=d.capex)
    p(c, "om", "O&M per year", "=capex*om_fraction", "euro/year", "", "#,##0",
      check=d.om_per_year)
    p(c, "elec_saved", "Electricity saved", "=cool_sol/comp_cop-cool_sol*parasitic", "kWh/year",
      "The conventional chiller's electricity, less the pumps and fans of this one", "#,##0")
    p(c, "savings", "Savings per year", "=heat_sol*heat_price+elec_saved*elec_price",
      "euro/year", "", "#,##0", check=d.savings_per_year)
    p(c, "useful", "Useful solar energy", "=heat_sol+cool_sol", "kWh/year", "Heat plus cold",
      "#,##0", check=d.solar_useful_kwh)
    p(c, "net", "Net benefit per year", "=savings-om", "euro/year", "", "#,##0")
    p(c, "lcoh", "Cost of the solar energy",
      f"=IF(useful>0,(capex*{_crf('disc_rate', 'life_years')}+om)/useful,\"none\")", "euro/kWh",
      "(CAPEX x CRF + O&M) / useful energy", "0.0000", check=d.lcoh)
    p(c, "payback", "Simple payback", '=IF(net>0,capex/net,"never")', "years", "", "0.0",
      check=d.payback_years)
    p(c, "npv_euro", "Net present value",
      f"=-capex+net*{_annuity('disc_rate', 'life_years')}", "euro", "", "#,##0", check=d.npv)
    p(c, "co2_t_year", "CO2 avoided", "=(heat_sol*heat_ef+elec_saved*grid_ef)/1000", "t CO2/year", "",
      "#,##0.00", check=d.co2_avoided_t)
    p(c, "yield_m2", "Useful energy per m2", "=useful/area", "kWh/m2/year", "", "0.0",
      check=d.yield_kwh_m2)

    ds = book.wb.create_sheet("Draw shape")
    ds.append(["Hour", "Share of the daily hot-water draw (weights)"])
    ds.cell(1, 1).font = _BOLD
    for hour, w in enumerate(st.DAILY_DRAW_SHAPE):
        ds.append([hour, w])
    ds.append(["Sum", "=SUM(B2:B25)"])

    h = book.sheet("Hourly", tuple([7] * 24))
    _table_header(h, [
        "Day", "Hour", "G (W/m2)", "Air (degC)", "Mains (degC)", "Hot-water demand (kWh)",
        "Cooling shape", "Cooling demand (kWh)", "Store T start", "Collector heat (kWh)",
        "T after collector", "Dumped (kWh)", "T after dump", "Stagnation", "Draw share",
        "Heat from sun (kWh)", "T after draw", "Chiller COP", "Chiller heat (kWh)",
        "Cold made (kWh)", "Too cold", "T after chiller", "Standing loss (kWh)",
        "Store T end"])
    for t in range(_HOURS):
        r = t + 2
        day, hour = divmod(t, 24)
        prev = "t_start" if t == 0 else f"X{r - 1}"
        h.append([
            day, hour, float(weather.g_w_m2[t]), float(weather.t_amb_c[t]),
            f"=MAX(4,t_mean-0.5*t_amp*COS(2*PI()*(A{r}-50)/365))",
            f"=dhw_l*INDEX('Draw shape'!$B$2:$B$25,B{r}+1)/'Draw shape'!$B$26"
            f"*{st.WATER_KWH_PER_L_K}*MAX(0,t_use-E{r})",
            f"=IF(AND(B{r}>=8,B{r}<20),MAX(0,D{r}-18),0)",
            f"=IF(AND(cool_kwh>0,raw_total>0),cool_kwh*G{r}/raw_total,0)",
            f"={prev}",
            f"=IF(AND(area>0,C{r}>0),MAX(0,coll_eta0-loss_k1*(I{r}+flow_dt/2-D{r})/C{r}"
            f"-loss_k2*(I{r}+flow_dt/2-D{r})^2/C{r})*C{r}*area/1000*(1-loop_loss),0)",
            f"=I{r}+J{r}/cap",
            f"=MAX(0,K{r}-{st.T_STAGNATION})*cap",
            f"=MIN(K{r},{st.T_STAGNATION})",
            f"=IF(K{r}>{st.T_STAGNATION},1,0)",
            f"=IF(F{r}>0,MIN(1,MAX(0,(M{r}-E{r})/MAX(1E-9,t_use-E{r}))),0)",
            f"=F{r}*O{r}",
            f"=M{r}-P{r}/cap",
            f"=IF(Q{r}<t_min,0,IF(Q{r}>=t_nom,cop_nom,"
            f"cop_nom*(0.45+0.55*(Q{r}-t_min)/MAX(1E-9,t_nom-t_min))))",
            f"=IF(AND(H{r}>0,R{r}>0,chiller_kw>0),MIN(H{r}/R{r},heat_cap,"
            f"MAX(0,(Q{r}-t_min)*cap)),0)",
            f"=S{r}*R{r}",
            f"=IF(AND(H{r}>0,NOT(AND(R{r}>0,chiller_kw>0))),1,0)",
            f"=Q{r}-S{r}/cap",
            f"=ua*MAX(0,V{r}-t_room)/1000",
            f"=MAX(V{r}-W{r}/cap,4)"])
    h.freeze_panes = "C2"
    book.finish_checks()
    return book.wb


# --- hydrogen ------------------------------------------------------------------------------
def _grid_by_hour(inp: HydrogenInputs, prof: HydrogenProfiles, res: list[float],
                  p_el: float) -> list[float] | None:
    """Grid energy of each hour when the cheapest hours are bought first (else None)."""
    if not (inp.use_grid and inp.cheapest_hours_first and prof.prices is not None):
        return None
    used = [min(r, p_el) for r in res]
    gap = [p_el - u for u in used]
    left = inp.grid_hours_cap * p_el
    grid = [0.0] * _HOURS
    for t in sorted(range(_HOURS), key=lambda i: prof.prices[i]):
        take = min(gap[t], left)
        if take <= 0:
            if left <= 0:
                break
            continue
        grid[t] = take
        left -= take
    return grid


def build_hydrogen_workbook(inp: HydrogenInputs, prof: HydrogenProfiles, d: HydrogenDesign
                            ) -> Workbook:
    book = _Book()
    p = book.put
    p_el, p_pv, p_wind = d.electrolyser_kw, d.pv_kwp, d.wind_kw
    pv_unit = list(prof.pv_unit) if prof.pv_unit and inp.use_pv else [0.0] * _HOURS
    wd_unit = list(prof.wind_unit) if prof.wind_unit and inp.use_wind else [0.0] * _HOURS
    res = [p_pv * a + p_wind * b for a, b in zip(pv_unit, wd_unit, strict=True)]
    fixed_grid = _grid_by_hour(inp, prof, res, p_el)
    prices = list(prof.prices) if prof.prices is not None else None
    _about(book, "RESize - green hydrogen as a spreadsheet", _COMMON_ABOUT + [
        "The sheet 'Hourly' rebuilds the year: renewable power up to the electrolyser's rating, "
        "grid top-up within its hour limit, minimum load, hydrogen made (flat or part-load "
        "efficiency), and the swing of the tank. 'Ageing' averages the stack's drift over its life.",
        "The renewable profiles (per kWp and per kW) are values. The sizes are those found by "
        "RESize; in 'demand' mode the program scales a 1 kW run to meet the hydrogen demand, "
        "here you can type another size instead.",
        "Not reproduced: the cheapest-hours-first purchase order (the hours chosen are pasted as "
        "values when that option is on).",
    ])
    ws = book.sheet("Inputs")
    _title(ws, "Inputs (yellow cells can be changed)")
    book.heading(ws, "Plant sizes")
    p(ws, "p_el", "Electrolyser", p_el, "kW", "", "#,##0.00", True)
    p(ws, "pv_kwp", "PV power", p_pv, "kWp", "0 if not used", "#,##0.00", True)
    p(ws, "wind_kw", "Wind power", p_wind, "kW", "0 if not used", "#,##0.00", True)
    p(ws, "use_grid", "Grid top-up (1 = yes, 0 = no)", 1 if inp.use_grid else 0, "", "", "0",
      True)
    p(ws, "grid_cap_h", "Grid limit", inp.grid_hours_cap, "full-load h/year", "", "#,##0", True)
    book.heading(ws, "Electrolyser and storage")
    p(ws, "kwh_stack", "Electricity per kg, stack and auxiliaries", inp.kwh_per_kg_stack,
      "kWh/kg", "At nominal load", "0.0", True)
    p(ws, "min_load", "Minimum load", inp.min_load, "-", "", "0.00", True)
    p(ws, "part_load", "Part-load efficiency model (1 = on)", 1 if inp.part_load_model else 0, "",
      "", "0", True)
    p(ws, "aux_share", "Auxiliaries, share of nominal power", inp.aux_share, "-", "", "0.00",
      True)
    p(ws, "stack_gain", "Stack gain at very low load", inp.stack_gain, "-", "", "0.00", True)
    p(ws, "degr", "Extra consumption per 1000 operating hours", inp.degradation_per_1000h, "-",
      "0.002 = 0.2 %", "0.0000", True)
    p(ws, "stack_hours", "Stack life", inp.stack_hours, "h", "", "#,##0", True)
    p(ws, "life_el", "Electrolyser life", inp.electrolyser_life_years, "years", "", "0.0", True)
    p(ws, "outlet_bar", "Electrolyser outlet pressure", inp.outlet_bar, "bar", "", "0.0", True)
    p(ws, "storage_bar", "Storage pressure", inp.storage_bar, "bar", "", "0", True)
    p(ws, "storage_days", "Storage autonomy ('days' mode)", inp.storage_days, "days", "", "0.0",
      True)
    p(ws, "hourly_tank", "Size the tank on the hourly swing (1) or on days (0)",
      1 if inp.storage_mode == "hourly" else 0, "", "", "0", True)
    p(ws, "water_l", "Water per kg of hydrogen", inp.water_l_per_kg, "L/kg", "", "0.0", True)
    book.heading(ws, "Costs")
    p(ws, "el_capex", "Electrolyser cost", inp.electrolyser_capex_per_kw, "euro/kW", "", "#,##0",
      True)
    p(ws, "el_om", "Electrolyser O&M, share of cost per year", inp.electrolyser_om_fraction, "-",
      "", "0.000", True)
    p(ws, "stack_cost_frac", "Stack replacement, share of the electrolyser cost",
      inp.stack_cost_fraction, "-", "", "0.00", True)
    p(ws, "tank_capex", "Storage cost", inp.storage_capex_per_kg, "euro/kg", "", "#,##0", True)
    p(ws, "water_price", "Water price", inp.water_price_per_m3, "euro/m3", "", "0.00", True)
    p(ws, "pv_capex", "PV cost", inp.pv_capex_per_kwp, "euro/kWp", "", "#,##0", True)
    p(ws, "pv_om", "PV O&M", inp.pv_om_per_kwp, "euro/kWp/year", "", "0.0", True)
    p(ws, "wind_capex", "Wind cost", inp.wind_capex_per_kw, "euro/kW", "", "#,##0", True)
    p(ws, "wind_om", "Wind O&M", inp.wind_om_per_kw, "euro/kW/year", "", "0.0", True)
    p(ws, "res_life", "Renewable plant life", inp.res_life_years, "years", "", "0.0", True)
    p(ws, "disc_rate", "Discount rate", inp.discount_rate, "-", "", "0.0%", True)
    p(ws, "grid_price", "Grid price (flat)", inp.grid_price_per_kwh, "euro/kWh",
          "Hourly prices are used when the profile has them.", "0.000", True)
    book.heading(ws, "Emissions")
    p(ws, "grid_ef", "Grid", inp.grid_ef, "kg CO2/kWh", "", "0.000", True)
    p(ws, "pv_ef", "PV", inp.pv_ef, "kg CO2/kWh", "", "0.000", True)
    p(ws, "wind_ef", "Wind", inp.wind_ef, "kg CO2/kWh", "", "0.000", True)
    p(ws, "grey_h2", "Grey hydrogen", inp.grey_h2_kg_co2_per_kg, "kg CO2/kg", "", "0.0", True)

    c = book.sheet("Calculation")
    _title(c, "Calculation (every cell is a formula)")
    n = _HOURS + 1
    book.heading(c, "1. Per kilogram")
    p(c, "compression", "Compression work",
      f"={h2.H2_GAS_CONSTANT_J_PER_KG_K}*298*LN(MAX(storage_bar,outlet_bar)/outlet_bar)"
      "/3.6E6/0.6", "kWh/kg", "R T ln(p2/p1) / isothermal efficiency 0.6", "0.000",
      check=d.compression_kwh_per_kg)
    p(c, "kwh_total", "Nominal electricity per kg", "=kwh_stack+compression", "kWh/kg", "",
      "0.00", check=d.nominal_kwh_per_kg_total)
    p(c, "stack_kwh", "Stack and auxiliaries only", "=MAX(kwh_total-compression,1E-9)", "kWh/kg",
      "", "0.00")
    book.heading(c, "2. The year (sums of the 'Hourly' sheet)")
    p(c, "res_avail", "Renewable energy available", f"=SUM(Hourly!D2:D{n})", "kWh", "",
      "#,##0", check=d.res_available_kwh)
    p(c, "grid_potential", "Gap to the electrolyser's rating", f"=SUM(Hourly!F2:F{n})", "kWh",
      "", "#,##0")
    p(c, "res_used", "Renewable energy used", f"=SUM(Hourly!L2:L{n})", "kWh", "", "#,##0",
      check=d.res_used_kwh)
    p(c, "grid_kwh", "Grid energy used", f"=SUM(Hourly!M2:M{n})", "kWh", "", "#,##0",
      check=d.grid_kwh)
    p(c, "electricity", "Electricity used", "=res_used+grid_kwh", "kWh", "", "#,##0",
      check=d.electricity_kwh)
    p(c, "hours_on", "Hours of operation", f"=SUM(Hourly!I2:I{n})", "h", "", "0",
      check=d.hours_on)
    p(c, "kg_year1", "Hydrogen, first year", f"=SUM(Hourly!K2:K{n})", "kg", "New stack",
      "#,##0", check=d.hydrogen_kg_year1)
    p(c, "kg", "Hydrogen per year (average over the life)", "=AVERAGE(Ageing!B2:B61)", "kg",
      "Mean over the years, with stack ageing", "#,##0", check=d.hydrogen_kg)
    p(c, "kg_hour_mean", "Mean hydrogen per hour", f"=SUM(Hourly!K2:K{n})/{_HOURS}", "kg/h",
      "The constant offtake the tank is sized against", "0.0000")
    p(c, "pv_used", "PV energy used", f"=SUM(Hourly!N2:N{n})", "kWh", "", "#,##0")
    p(c, "wind_used", "Wind energy used", f"=SUM(Hourly!O2:O{n})", "kWh", "", "#,##0")
    p(c, "flh", "Full-load hours", "=electricity/p_el", "h/year", "", "#,##0",
      check=d.full_load_hours)
    p(c, "kwh_per_kg", "Electricity per kg, as operated", "=electricity/kg", "kWh/kg", "",
      "0.00", check=d.kwh_per_kg_total)
    p(c, "eff", "Efficiency (LHV)", f"={h2.H2_LHV_KWH_PER_KG}/kwh_per_kg", "-",
      "33.33 kWh/kg / electricity per kg", "0.0%", check=d.efficiency_lhv)
    p(c, "curtailed", "Renewable energy not used", "=res_avail-res_used", "kWh", "", "#,##0",
      check=d.curtailed_kwh)
    p(c, "stack_life", "Stack life in years", "=stack_hours/flh", "years", "", "0.0",
      check=d.stack_life_years)
    book.heading(c, "3. Water, oxygen and storage")
    p(c, "water_m3", "Water", "=kg*water_l/1000", "m3/year", "", "#,##0.0", check=d.water_m3)
    p(c, "oxygen_t", "Oxygen", f"=kg*{h2.OXYGEN_KG_PER_KG}/1000", "t/year", "", "#,##0.0",
      check=d.oxygen_t)
    p(c, "tank_hi", "Highest point of the tank balance", f"=MAX(0,MAX(Hourly!Q2:Q{n}))", "kg",
      "", "#,##0.0")
    p(c, "tank_lo", "Lowest point of the tank balance", f"=MIN(0,MIN(Hourly!Q2:Q{n}))", "kg",
      "", "#,##0.0")
    p(c, "storage_kg", "Hydrogen stored", "=IF(hourly_tank=1,tank_hi-tank_lo,kg/365*storage_days)",
      "kg", "Hourly swing, or days of autonomy", "#,##0.0", check=d.storage_kg)
    p(c, "density", "Density at the storage pressure",
      f"=storage_bar*100000/((1+0.0006*storage_bar)*{h2.H2_GAS_CONSTANT_J_PER_KG_K}*288)",
      "kg/m3", "Real-gas correction Z = 1 + 0.0006 p", "0.00")
    p(c, "storage_m3", "Tank volume", "=storage_kg/density", "m3", "", "#,##0.00",
      check=d.storage_m3)
    book.heading(c, "4. Cost of the hydrogen (annualised)")
    p(c, "crf_el", "CRF, electrolyser", f"={_crf('disc_rate', 'life_el')}", "1/year", "", "0.0000")
    p(c, "crf_res", "CRF, renewables", f"={_crf('disc_rate', 'res_life')}", "1/year", "", "0.0000")
    p(c, "capex_el", "Electrolyser cost", "=p_el*el_capex", "euro", "", "#,##0")
    p(c, "cost_el", "Electrolyser, per year", "=capex_el*(crf_el+el_om)", "euro/year", "",
      "#,##0", check=d.cost_electrolyser)
    p(c, "cost_stack", "Stack replacement, per year",
      "=capex_el*stack_cost_frac*flh/stack_hours", "euro/year", "", "#,##0", check=d.cost_stack)
    p(c, "cost_storage", "Storage, per year", "=storage_kg*tank_capex*crf_el", "euro/year", "",
      "#,##0", check=d.cost_storage)
    p(c, "cost_res", "Renewable plants, per year",
      "=pv_kwp*(pv_capex*crf_res+pv_om)+wind_kw*(wind_capex*crf_res+wind_om)", "euro/year", "",
      "#,##0", check=d.cost_res)
    grid_cost = (f"=SUM(Hourly!P2:P{n})" if prices is not None else "=grid_kwh*grid_price")
    p(c, "cost_grid", "Grid electricity, per year", grid_cost, "euro/year", "", "#,##0",
      check=d.cost_grid)
    p(c, "cost_water", "Water, per year", "=water_m3*water_price", "euro/year", "", "#,##0",
      check=d.cost_water)
    p(c, "total", "Total per year",
      "=cost_el+cost_stack+cost_storage+cost_res+cost_grid+cost_water", "euro/year", "",
      "#,##0")
    p(c, "lcoh", "Cost of hydrogen (LCOH)", "=total/kg", "euro/kg", "Total per year / kg",
      "0.000", check=d.lcoh)
    book.heading(c, "5. Carbon")
    p(c, "co2_kg", "Emissions per kg of hydrogen",
      "=(grid_kwh*grid_ef+pv_used*pv_ef+wind_used*wind_ef)/kg", "kg CO2/kg", "", "0.000",
      check=d.co2_kg_per_kg)
    p(c, "co2_saved", "CO2 saved against grey hydrogen", "=(grey_h2-co2_kg)*kg/1000",
      "t CO2/year", "", "#,##0.0", check=d.co2_saved_t_per_year)

    ag = book.sheet("Ageing", (10, 22))
    _table_header(ag, ["Year", "Hydrogen (kg)"])
    life_years = max(1, round(inp.electrolyser_life_years))
    for y in range(1, 61):
        r = y + 1
        ag.append([y, f'=IF(A{r}>MAX(1,ROUND(life_el,0)),"",kg_year1/(1+degr*'
                      f'IF(stack_hours>0,MOD((A{r}-0.5)*hours_on,stack_hours),(A{r}-0.5)*hours_on)'
                      f'/1000*stack_kwh/kwh_total))'])
    ag.append([])
    ag.append([f"The stack drifts with the hours it has run; a new one is fitted after the stack "
               f"life. Years counted: {life_years} (from the electrolyser life)."])

    hr = book.sheet("Hourly", tuple([11] * 17))
    _table_header(hr, [
        "Hour", "PV per kWp", "Wind per kW", "Renewable (kW)", "Used (kW)", "Gap to rating",
        "Grid (kW)", "Total (kW)", "On", "Load share", "Hydrogen (kg)", "Renewable used on",
        "Grid used on", "PV used", "Wind used", "Grid cost (euro)", "Tank balance (kg)"])
    scale_ref = "MIN(1,grid_cap_h*p_el/grid_potential)"
    for t in range(_HOURS):
        r = t + 2
        grid = (fixed_grid[t] if fixed_grid is not None
                else f"=IF(AND(use_grid=1,grid_potential>0),{scale_ref}*F{r},0)")
        price = f"=M{r}*{prices[t]!r}" if prices is not None else None
        prev_tank = "0" if t == 0 else f"Q{r - 1}"
        hr.append([
            t, pv_unit[t], wd_unit[t], f"=pv_kwp*B{r}+wind_kw*C{r}", f"=MIN(D{r},p_el)",
            f"=IF(use_grid=1,p_el-E{r},0)", grid, f"=E{r}+G{r}",
            f"=IF(AND(H{r}>=min_load*p_el,H{r}>0),1,0)", f"=H{r}/p_el",
            f"=IF(I{r}=1,IF(part_load=1,H{r}/(stack_kwh*((1-aux_share)*(1-stack_gain*(1-MIN(J{r},1)))"
            f"+aux_share/MIN(J{r},1))+compression),H{r}/kwh_total),0)",
            f"=I{r}*E{r}", f"=I{r}*G{r}",
            f"=I{r}*IF(D{r}>0,pv_kwp*B{r}*E{r}/D{r},0)",
            f"=I{r}*IF(D{r}>0,wind_kw*C{r}*E{r}/D{r},0)",
            price if price else 0,
            f"={prev_tank}+K{r}-kg_hour_mean"])
    hr.freeze_panes = "B2"
    book.finish_checks()
    return book.wb


__all__ = ["build_battery_workbook", "build_biomass_workbook", "build_hydrogen_workbook",
           "build_solar_workbook"]
